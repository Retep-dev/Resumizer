"""1–4: accurate parsing, title/employer preservation, no invented experience."""
import asyncio
import io
import json

import docx
import pytest

from app.agents.resume_parser_agent import ground_resume, parse_resume_with_ai, rule_based_parse
from app.schemas.resume import ResumeSchema
from app.services.document_parser import DocumentError, extract_text_from_file
from tests.conftest import FRONTEND_RESUME, PRIYA_RESUME, FakeClient, priya_parse_json


def run(coro):
    return asyncio.run(coro)


# ── Text extraction ──

def test_txt_and_md_extraction_keeps_content():
    assert "Halcyon Freight" in extract_text_from_file(PRIYA_RESUME.encode(), "cv.txt")
    assert "Kitebird" in extract_text_from_file(FRONTEND_RESUME.encode(), "cv.md")


def test_docx_extraction_includes_tables():
    d = docx.Document()
    d.add_paragraph("Priya Raman")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Senior Data Engineer - Halcyon Freight"
    table.cell(0, 1).text = "Mar 2021 - Present"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text_from_file(buf.getvalue(), "cv.docx")
    assert "Priya Raman" in text and "Halcyon Freight" in text and "Mar 2021" in text


def test_pdf_extraction_reads_text_layer():
    matplotlib = pytest.importorskip("matplotlib")
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(8.5, 11))
    for i, line in enumerate(PRIYA_RESUME.splitlines()[:12]):
        fig.text(0.08, 0.95 - i * 0.03, line.replace("•", "-"), fontsize=10)
    buf = io.BytesIO()
    fig.savefig(buf, format="pdf")
    plt.close(fig)
    text = extract_text_from_file(buf.getvalue(), "cv.pdf")
    assert "Halcyon Freight" in text and "Senior Data Engineer" in text


def test_legacy_doc_and_unknown_types_are_rejected_clearly():
    with pytest.raises(DocumentError) as e:
        extract_text_from_file(b"\xd0\xcf\x11\xe0", "cv.doc")
    assert e.value.code == "validation_failed" and "DOCX" in e.value.message
    with pytest.raises(DocumentError):
        extract_text_from_file(b"abc", "photo.png")


# ── AI parsing is grounded in the resume ──

def test_ai_parse_preserves_titles_employers_and_dates():
    resume, warnings, model = run(parse_resume_with_ai(FakeClient(parse=priya_parse_json()), PRIYA_RESUME))
    jobs = [(j.job_title, j.company, j.start_date, j.end_date) for j in resume.work_experience]
    assert jobs == [("Senior Data Engineer", "Halcyon Freight", "Mar 2021", "Present"),
                    ("Data Analyst", "Brightwell Insurance", "Jul 2018", "Feb 2021")]
    assert resume.full_name == "Priya Raman"
    assert len(resume.work_experience[0].bullet_points) == 3
    assert warnings == []


def test_grounding_removes_hallucinated_role_skills_and_contact():
    bad = json.loads(priya_parse_json(
        email="candidate@example.com",
        technical_skills=["Python", "Kubernetes", "System Design"],
    ))
    bad["work_experience"].append({"company": "Engineering Experience", "job_title": "Software Developer",
                                   "start_date": "", "end_date": "Present", "bullet_points": ["Architected cloud systems."]})
    bad["education"].append({"institution": "University", "degree": "Bachelor of Science in Computer Science", "graduation_year": "N/A"})
    bad["work_experience"][0]["bullet_points"].append("Led a team of 12 engineers to migrate to Kubernetes.")
    resume, warnings = ground_resume(ResumeSchema.model_validate(bad), PRIYA_RESUME)

    assert [j.company for j in resume.work_experience] == ["Halcyon Freight", "Brightwell Insurance"]
    assert "Kubernetes" not in resume.technical_skills and "System Design" not in resume.technical_skills
    assert resume.email is None  # placeholder email isn't in the resume
    assert all("12 engineers" not in b for j in resume.work_experience for b in j.bullet_points)
    assert [e.institution for e in resume.education] == ["University of Leeds"]
    assert any("Removed a role" in w for w in warnings)


def test_grounding_blanks_a_wrong_name():
    resume, warnings = ground_resume(ResumeSchema.model_validate(json.loads(priya_parse_json(full_name="Jane Doe"))), PRIYA_RESUME)
    assert resume.full_name == ""
    assert warnings


# ── Rule-based parser (AI unavailable) only copies resume text ──

@pytest.mark.parametrize("text,expected", [
    (PRIYA_RESUME, [("Senior Data Engineer", "Halcyon Freight"), ("Data Analyst", "Brightwell Insurance")]),
    (FRONTEND_RESUME, [("Frontend Developer", "Kitebird"), ("Junior Web Developer", "Palmline Media")]),
])
def test_rule_based_parser_keeps_real_titles_and_employers(text, expected):
    resume, warnings = rule_based_parse(text)
    assert [(j.job_title, j.company) for j in resume.work_experience] == expected
    assert all(j.start_date and j.end_date for j in resume.work_experience)
    assert warnings  # always labelled as pattern-based


def test_rule_based_parser_never_invents_experience():
    resume, _ = rule_based_parse(PRIYA_RESUME)
    blob = resume.model_dump_json()
    for invented in ("Engineering Experience", "Software Developer", "Computer Science", "candidate@example.com", "Software Project"):
        assert invented not in blob
    # every bullet is copied from the resume
    for job in resume.work_experience:
        for b in job.bullet_points:
            assert b in PRIYA_RESUME
    assert resume.full_name == "Priya Raman"
    assert resume.education[0].institution == "University of Leeds"


def test_rule_based_parser_with_no_structure_returns_empty_not_fake():
    resume, warnings = rule_based_parse("just some words without any structure at all " * 5)
    assert resume.work_experience == [] and resume.education == []
    assert any("No roles" in w for w in warnings)
